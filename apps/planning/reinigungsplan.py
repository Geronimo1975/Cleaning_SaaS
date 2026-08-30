"""
Export Reinigungsplan (XLSX).

Concurența vinde un șablon Excel gol cu patru coloane: Bereich, Leistung,
Beschreibung, Reinigungsfrequenz. Aici exportăm aceleași patru coloane, dar
completate din datele reale ale obiectivului — plus durata pe vizită și
minutele săptămânale, care sunt exact ce nu poate produce un șablon gol.

Documentul e pentru client și pentru lucrător, deci e în germană.
"""

from io import BytesIO

from .models import Zone

WEEKDAY_DE = {1: "Mo", 2: "Di", 3: "Mi", 4: "Do", 5: "Fr", 6: "Sa", 7: "So"}

HEADERS = [
    ("Bereich / Ort", 34),
    ("Leistung", 30),
    ("Beschreibung", 46),
    ("Reinigungsfrequenz", 22),
    ("Dauer je Einsatz", 18),
    ("Wochentage", 18),
    ("Minuten / Woche", 16),
]


def frequency_label(visits_per_week: int) -> str:
    if visits_per_week >= 7:
        return "täglich"
    if visits_per_week == 1:
        return "1× wöchentlich"
    return f"{visits_per_week}× wöchentlich"


def zone_rows(site) -> list[list]:
    zones = (
        Zone.objects.filter(site=site, is_active=True)
        .select_related("category")
        .order_by("name")
    )
    return [
        [
            f"{site.name} · {z.name}",
            z.category.name_de or z.category.name,
            z.description,
            frequency_label(z.visits_per_week),
            f"{z.minutes_per_visit} min",
            ", ".join(WEEKDAY_DE[d] for d in z.weekdays),
            z.weekly_minutes,
        ]
        for z in zones
    ]


def build_workbook(site) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reinigungsplan"

    ws["A1"] = f"Reinigungsplan — {site.name}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"{site.street}, {site.postal_code} {site.locality.name}".strip(", ")
    ws["A2"].font = Font(size=10, color="666666")

    header_row = 4
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    for col, (label, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    rows = zone_rows(site)
    for i, row in enumerate(rows, start=header_row + 1):
        for col, value in enumerate(row, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.alignment = Alignment(wrap_text=col == 3, vertical="top")

    total_row = header_row + len(rows) + 1
    ws.cell(row=total_row, column=1, value="Summe").font = Font(bold=True)
    total = ws.cell(row=total_row, column=7, value=sum(r[6] for r in rows))
    total.font = Font(bold=True)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def filename(site) -> str:
    from django.utils.text import slugify

    return f"reinigungsplan-{slugify(site.name)}.xlsx"
