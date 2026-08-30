"""
Export Reinigungsplan.

Contractul: cele patru coloane pe care concurența le livrează goale
(Bereich, Leistung, Beschreibung, Reinigungsfrequenz) ies completate din
datele obiectivului, plus durata și minutele săptămânale.
"""

from datetime import date

import pytest

from apps.planning.models import Site, Zone
from apps.planning.reinigungsplan import build_workbook, frequency_label, zone_rows

MONDAY = date(2026, 10, 5)


@pytest.fixture
def site(db, employer, norderstedt):
    return Site.objects.create(
        account=employer, name="Haus Nord", locality=norderstedt,
        street="Ulzburger Str. 40", postal_code="22846",
    )


@pytest.fixture
def zone(db, site, category):
    def _make(name="Treppenhaus A", minutes=45, visits=2, weekdays="2,4", description=""):
        return Zone.objects.create(
            site=site, name=name, category=category, minutes_per_visit=minutes,
            visits_per_week=visits, allowed_weekdays=weekdays, description=description,
        )

    return _make


@pytest.mark.django_db
class TestRows:
    def test_coloanele_vin_din_date_reale(self, site, zone):
        zone(description="Handläufe feucht wischen, Fenstersimse abstauben.")
        (row,) = zone_rows(site)
        bereich, leistung, beschreibung, frequenz, dauer, tage, wochenminuten = row
        assert bereich == "Haus Nord · Treppenhaus A"
        assert leistung  # denumirea DE a categoriei
        assert beschreibung.startswith("Handläufe")
        assert frequenz == "2× wöchentlich"
        assert dauer == "45 min"
        assert tage == "Di, Do"
        assert wochenminuten == 90

    def test_zonele_inactive_nu_apar(self, site, zone):
        z = zone()
        z.is_active = False
        z.save(update_fields=["is_active"])
        assert zone_rows(site) == []

    @pytest.mark.parametrize(
        "visits,expected",
        [(1, "1× wöchentlich"), (3, "3× wöchentlich"), (7, "täglich")],
    )
    def test_eticheta_de_frecventa(self, visits, expected):
        assert frequency_label(visits) == expected


@pytest.mark.django_db
class TestWorkbook:
    def test_fisierul_contine_antetul_si_totalul(self, site, zone):
        from openpyxl import load_workbook

        zone(name="Treppenhaus A", minutes=45, visits=2)
        zone(name="Tiefgarage", minutes=90, visits=1)
        ws = load_workbook(build_workbook(site)).active

        assert ws["A1"].value == "Reinigungsplan — Haus Nord"
        assert [c.value for c in ws[4]][:4] == [
            "Bereich / Ort", "Leistung", "Beschreibung", "Reinigungsfrequenz"
        ]
        # 2 zone -> rândurile 5 și 6, total pe 7
        assert ws.cell(row=7, column=1).value == "Summe"
        assert ws.cell(row=7, column=7).value == 90 + 90


@pytest.mark.django_db
class TestView:
    def test_proprietarul_descarca_xlsx(self, client, employer, site, zone):
        zone()
        client.force_login(employer.user)
        r = client.get(f"/dashboard/obiective/{site.pk}/reinigungsplan.xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r["Content-Type"]
        assert "reinigungsplan-haus-nord.xlsx" in r["Content-Disposition"]

    def test_alt_cont_nu_are_acces(self, client, worker, site, zone):
        zone()
        client.force_login(worker.user)
        r = client.get(f"/dashboard/obiective/{site.pk}/reinigungsplan.xlsx")
        assert r.status_code in (403, 404)
